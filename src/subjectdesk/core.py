"""Assignment discovery, persistent state, model calls, and document output."""

from __future__ import annotations

import fcntl
import html
import json
import os
import re
import shutil
import sqlite3
import uuid
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Any, cast

import yaml
from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class SubjectdeskError(Exception):
    """An assignment could not be processed."""


INPUT_SUFFIXES = {".md", ".txt", ".docx", ".pdf"}
OUTPUT_FORMATS = {"docx", "pdf", "md", "txt", "html"}
MODELS = {"haiku": "claude-haiku-4-5", "sonnet": "claude-sonnet-5-5"}
RULES = """Write plainly and naturally. Do not use em dashes or emojis. Do not roleplay,
adopt a character, or describe yourself as an AI. State uncertainty honestly. Research
factual claims with the web search tool. Cite source URLs close to the claims they support.
Never invent sources, facts, research results, or completed work. Follow the assignment's
format and content instructions where they do not conflict with these rules.
"""
EMOJI = re.compile("[\U0001F000-\U0001FAFF\u2600-\u27BF]")
CITE_TAG = re.compile(r"</?cite(?:\s+[^>]*)?>")


@dataclass(frozen=True)
class Assignment:
    title: str
    text: str
    formats: tuple[str, ...]
    model: str


def state_dir(root: Path) -> Path:
    return root / ".subjectdesk"


@contextmanager
def workspace_lock(root: Path) -> Iterator[None]:
    lock_path = state_dir(root) / "watch.lock"
    try:
        with lock_path.open("a+") as handle:
            try:
                fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError as exc:
                raise SubjectdeskError(f"another worker is using {root}") from exc
            try:
                yield
            finally:
                fcntl.flock(handle, fcntl.LOCK_UN)
    except FileNotFoundError as exc:
        raise SubjectdeskError(f"not initialized: {root}; run subjectdesk init first") from exc


def requeue_interrupted(root: Path) -> int:
    with connect(root) as db:
        updated = db.execute(
            "UPDATE jobs SET status='queued', error=NULL, updated_at=CURRENT_TIMESTAMP "
            "WHERE status='running'"
        )
        return updated.rowcount


def init(root: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    folder = state_dir(root)
    folder.mkdir(exist_ok=True)
    rules = folder / "rules.txt"
    if not rules.exists():
        rules.write_text(RULES, encoding="utf-8")
    with connect(root):
        pass


@contextmanager
def connect(root: Path) -> Iterator[sqlite3.Connection]:
    if not state_dir(root).is_dir():
        raise SubjectdeskError(f"not initialized: {root}; run subjectdesk init first")
    db = sqlite3.connect(state_dir(root) / "jobs.sqlite", timeout=30)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("""CREATE TABLE IF NOT EXISTS jobs (
        id TEXT PRIMARY KEY, subject TEXT NOT NULL, source TEXT NOT NULL UNIQUE,
        output_dir TEXT NOT NULL, status TEXT NOT NULL, model TEXT NOT NULL,
        formats TEXT NOT NULL, error TEXT, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    )""")
    db.commit()
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def read_input(path: Path) -> str:
    try:
        if path.suffix.lower() in {".md", ".txt"}:
            return path.read_text(encoding="utf-8")
        if path.suffix.lower() == ".docx":
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        if path.suffix.lower() == ".pdf":
            return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    except Exception as exc:
        raise SubjectdeskError(f"cannot read {path.name}: {exc}") from exc
    raise SubjectdeskError(f"unsupported assignment type: {path.suffix}")


def parse_assignment(path: Path, default_model: str = "haiku") -> Assignment:
    raw = read_input(path)
    metadata: dict[str, Any] = {}
    body = raw
    if raw.startswith("---\n"):
        _, front, body = raw.split("---", 2)
        try:
            metadata = yaml.safe_load(front) or {}
        except yaml.YAMLError as exc:
            raise SubjectdeskError(f"invalid front matter in {path.name}: {exc}") from exc
        if not isinstance(metadata, dict):
            raise SubjectdeskError("front matter must be a YAML mapping")
    else:
        match = re.search(r"(?im)^formats?\s*:\s*([^\n]+)$", body)
        if match:
            metadata["formats"] = match.group(1)
    if not body.strip():
        raise SubjectdeskError("assignment is empty")
    raw_formats = metadata.get("formats", ["docx"])
    if isinstance(raw_formats, str):
        formats = tuple(x.strip().lower().lstrip(".") for x in raw_formats.split(","))
    elif isinstance(raw_formats, list):
        formats = tuple(str(x).strip().lower().lstrip(".") for x in raw_formats)
    else:
        raise SubjectdeskError("formats must be a list or comma-separated string")
    if not formats or any(x not in OUTPUT_FORMATS for x in formats):
        raise SubjectdeskError(f"formats must use: {', '.join(sorted(OUTPUT_FORMATS))}")
    model = str(metadata.get("model", default_model)).lower()
    if model not in MODELS:
        raise SubjectdeskError("model must be haiku or sonnet")
    title = str(metadata.get("title", path.stem.replace("_", " ").replace("-", " ")))
    return Assignment(title, body.strip(), tuple(dict.fromkeys(formats)), model)


def discover(
    root: Path, default_model: str = "haiku", allowed_paths: set[str] | None = None
) -> list[str]:
    """Record new top-level subject files. Existing jobs are never requeued automatically."""
    added: list[str] = []
    with connect(root) as db:
        for subject in sorted(root.iterdir()):
            if not subject.is_dir() or subject.name.startswith("."):
                continue
            for source in sorted(subject.iterdir()):
                if not source.is_file() or source.name.startswith("."):
                    continue
                if source.suffix.lower() not in INPUT_SUFFIXES:
                    continue
                if allowed_paths is not None and str(source) not in allowed_paths:
                    continue
                if db.execute("SELECT 1 FROM jobs WHERE source=?", (str(source),)).fetchone():
                    continue
                name = source.stem
                output = subject / name
                index = 2
                while output.exists() or db.execute(
                    "SELECT 1 FROM jobs WHERE output_dir=?", (str(output),)
                ).fetchone():
                    output = subject / f"{name}-{index}"
                    index += 1
                job_id = uuid.uuid4().hex[:12]
                try:
                    assignment = parse_assignment(source, default_model)
                    status, error = "queued", None
                    model, formats = assignment.model, assignment.formats
                except SubjectdeskError as exc:
                    status, error = "failed", str(exc)
                    model, formats = default_model, ()
                db.execute(
                    "INSERT INTO jobs (id, subject, source, output_dir, status, model, formats, error) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (job_id, subject.name, str(source), str(output), status, model,
                     json.dumps(formats), error),
                )
                added.append(job_id)
    return added


def jobs(root: Path) -> list[sqlite3.Row]:
    with connect(root) as db:
        return db.execute("SELECT * FROM jobs ORDER BY created_at, id").fetchall()


def claim(root: Path) -> sqlite3.Row | None:
    with connect(root) as db:
        while True:
            row = db.execute("SELECT * FROM jobs WHERE status='queued' ORDER BY created_at, id LIMIT 1").fetchone()
            if row is None:
                return None
            updated = db.execute(
                "UPDATE jobs SET status='running', updated_at=CURRENT_TIMESTAMP "
                "WHERE id=? AND status='queued'", (row["id"],)
            )
            if updated.rowcount:
                return row


def set_status(root: Path, job_id: str, status: str, error: str | None = None) -> None:
    with connect(root) as db:
        db.execute(
            "UPDATE jobs SET status=?, error=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (status, error, job_id),
        )


def retry(root: Path, job_id: str) -> None:
    with connect(root) as db:
        row = db.execute("SELECT status FROM jobs WHERE id=?", (job_id,)).fetchone()
        if row is None:
            raise SubjectdeskError(f"unknown job: {job_id}")
        if row["status"] not in {"failed", "running"}:
            raise SubjectdeskError(f"job is {row['status']}; only failed or interrupted jobs can retry")
        db.execute(
            "UPDATE jobs SET status='queued', error=NULL, updated_at=CURRENT_TIMESTAMP WHERE id=?",
            (job_id,),
        )


def research(assignment: Assignment, rules: str) -> tuple[str, str]:
    """Ask Claude to research with Anthropic's server-side web search."""
    if not os.getenv("ANTHROPIC_API_KEY"):
        raise SubjectdeskError("ANTHROPIC_API_KEY is not set")
    import anthropic

    client = anthropic.Anthropic(timeout=180)
    searched = False
    found_results = False
    messages: list[dict[str, Any]] = [{"role": "user", "content": (
        f"Assignment title: {assignment.title}\n\n{assignment.text}\n\n"
        "Research the assignment using web search before writing. Return only a JSON object "
        "with string fields title and body_markdown. Include a Sources section with direct URLs "
        "for sources actually used. If you cannot verify a claim, state the uncertainty."
    )}]
    for _ in range(4):
        response = client.messages.create(
            model=MODELS[assignment.model], max_tokens=8192,
            system=rules,
            tools=[{"type": "web_search_20250305", "name": "web_search", "max_uses": 8}],
            messages=cast(Any, messages),
        )
        for block in response.content:
            if block.type == "server_tool_use" and block.name == "web_search":
                searched = True
            if block.type == "web_search_tool_result":
                if isinstance(block.content, list) and block.content:
                    found_results = True
                elif not isinstance(block.content, list):
                    raise SubjectdeskError("web search returned an error")
        if response.stop_reason == "pause_turn":
            messages.append({"role": "assistant", "content": response.content})
            continue
        if response.stop_reason == "max_tokens":
            raise SubjectdeskError("model output reached the token limit")
        if response.stop_reason != "end_turn":
            raise SubjectdeskError(f"model stopped unexpectedly: {response.stop_reason}")
        if not searched or not found_results:
            raise SubjectdeskError("model did not complete a web search with results")
        text = "\n".join(block.text for block in response.content if block.type == "text")
        return parse_model_output(text)
    raise SubjectdeskError("model did not finish after four API turns")


def parse_model_output(raw: str) -> tuple[str, str]:
    cleaned = raw.strip()
    if cleaned.startswith("```json"):
        cleaned = cleaned[7:].removesuffix("```")
    try:
        result = json.loads(cleaned.strip())
    except (json.JSONDecodeError, ValueError) as exc:
        raise SubjectdeskError("model did not return valid JSON") from exc
    if not isinstance(result, dict) or not isinstance(result.get("title"), str) or not isinstance(result.get("body_markdown"), str):
        raise SubjectdeskError("model output needs title and body_markdown strings")
    title = clean_style(result["title"].strip())
    body = clean_style(result["body_markdown"].strip())
    if not title or not body:
        raise SubjectdeskError("model returned an empty document")
    return title, body


def clean_style(value: str) -> str:
    return EMOJI.sub("", CITE_TAG.sub("", value.replace("\u2014", ", "))).strip()


def markdown_lines(body: str) -> list[tuple[str, str]]:
    lines: list[tuple[str, str]] = []
    for line in body.splitlines():
        stripped = line.strip()
        if stripped.startswith("### "):
            lines.append(("h3", stripped[4:]))
        elif stripped.startswith("## "):
            lines.append(("h2", stripped[3:]))
        elif stripped.startswith("# "):
            lines.append(("h1", stripped[2:]))
        elif stripped.startswith(("- ", "* ")):
            lines.append(("bullet", stripped[2:]))
        else:
            lines.append(("text", stripped))
    return lines


def render(title: str, body: str, formats: tuple[str, ...], directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    body = re.sub(r"^#\s+" + re.escape(title) + r"\s*\n+", "", body, count=1, flags=re.IGNORECASE)
    for fmt in formats:
        target = directory / f"work.{fmt}"
        if fmt == "md":
            target.write_text(f"# {title}\n\n{body}\n", encoding="utf-8")
        elif fmt == "txt":
            target.write_text(f"{title}\n\n{body}\n", encoding="utf-8")
        elif fmt == "html":
            parts = [f"<h1>{html.escape(title)}</h1>"]
            for kind, line in markdown_lines(body):
                if not line:
                    continue
                tag = "p" if kind in {"text", "bullet"} else kind
                prefix = "• " if kind == "bullet" else ""
                parts.append(f"<{tag}>{html.escape(prefix + line)}</{tag}>")
            target.write_text("<!doctype html><meta charset=\"utf-8\"><title>" +
                              html.escape(title) + "</title>\n" + "\n".join(parts), encoding="utf-8")
        elif fmt == "docx":
            doc = Document()
            doc.add_heading(title, 0)
            for kind, line in markdown_lines(body):
                if kind.startswith("h"):
                    doc.add_heading(line, level=int(kind[1]))
                elif kind == "bullet":
                    doc.add_paragraph(line, style="List Bullet")
                else:
                    doc.add_paragraph(line)
            doc.save(str(target))
        elif fmt == "pdf":
            _render_pdf(title, body, target)


def _render_pdf(title: str, body: str, target: Path) -> None:
    font_paths = (
        Path("/usr/share/fonts/TTF/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/liberation/LiberationSans-Regular.ttf"),
    )
    font = next((path for path in font_paths if path.exists()), None)
    if font is None:
        raise SubjectdeskError("PDF output requires a DejaVu or Liberation Sans TTF font")
    pdfmetrics.registerFont(TTFont("Subjectdesk", str(font)))
    styles = getSampleStyleSheet()
    for key in ("Title", "Heading1", "Heading2", "Heading3", "BodyText"):
        styles[key].fontName = "Subjectdesk"
    story: list[Any] = [Paragraph(html.escape(title), styles["Title"]), Spacer(1, 14)]
    for kind, line in markdown_lines(body):
        if not line:
            story.append(Spacer(1, 8))
            continue
        style = styles[{"h1": "Heading1", "h2": "Heading2", "h3": "Heading3"}.get(kind, "BodyText")]
        content = ("• " if kind == "bullet" else "") + line
        story.append(Paragraph(html.escape(content), style))
        story.append(Spacer(1, 5))
    SimpleDocTemplate(str(target), pagesize=A4).build(story)


def process_one(root: Path, researcher: Callable[[Assignment, str], tuple[str, str]] = research) -> str | None:
    row = claim(root)
    if row is None:
        return None
    source = Path(row["source"])
    output = Path(row["output_dir"])
    staging = state_dir(root) / "staging" / row["id"]
    try:
        if not source.is_file():
            raise SubjectdeskError(f"assignment file disappeared: {source}")
        assignment = parse_assignment(source, row["model"])
        rules = (state_dir(root) / "rules.txt").read_text(encoding="utf-8")
        title, body = researcher(assignment, RULES + "\n" + rules)
        title, body = clean_style(title), clean_style(body)
        if not title or not body:
            raise SubjectdeskError("researcher returned an empty document")
        render(title, body, assignment.formats, staging)
        if output.exists():
            raise SubjectdeskError(f"output folder already exists: {output}")
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, staging / source.name)
        staging.rename(output)
        set_status(root, row["id"], "done")
        try:
            source.unlink()
        except OSError:
            pass
    except Exception as exc:  # noqa: BLE001  Model, IO, and renderer failures all become job failures.
        set_status(root, row["id"], "failed", str(exc))
        shutil.rmtree(staging, ignore_errors=True)
    return row["id"]


def scan_stable(root: Path, seen: dict[str, tuple[int, int]], default_model: str) -> list[str]:
    """Wait for two scans with unchanged size and mtime before admitting a file."""
    current: dict[str, tuple[int, int]] = {}
    for subject in root.iterdir():
        if not subject.is_dir() or subject.name.startswith("."):
            continue
        for source in subject.iterdir():
            if source.is_file() and source.suffix.lower() in INPUT_SUFFIXES and not source.name.startswith("."):
                stat = source.stat()
                current[str(source)] = (stat.st_size, stat.st_mtime_ns)
    stable = {path for path, fingerprint in current.items() if seen.get(path) == fingerprint}
    seen.clear()
    seen.update(current)
    if not stable:
        return []
    return discover(root, default_model, stable)
